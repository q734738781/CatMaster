from __future__ import annotations

from pathlib import Path
import subprocess
import sys
from types import SimpleNamespace

from langgraph.store.memory import InMemoryStore
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import pytest

from catmaster.llm.config import AgentRuntimeConfig
from catmaster.specialists.runtime import SpecialistRunner, build_specialist_runner


RUNTIME_ROOT = Path(__file__).resolve().parents[1] / "third_party/easyslides"


def test_library_mount_and_shell_execution_share_complete_installation(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    library = tmp_path / "installed library"
    (library / "scripts").mkdir(parents=True)
    (library / "references").mkdir()
    (library / "references/example.md").write_text("Library reference content")
    (library / "scripts/example.py").write_text(
        "from pathlib import Path\n"
        "print((Path(__file__).resolve().parents[1] / 'references/example.md').read_text())\n"
        "Path('output.txt').write_text('workspace output')\n"
    )
    monkeypatch.setattr(SpecialistRunner, "_easyslides_root", staticmethod(lambda: library))
    workspace = tmp_path / "workspace"
    files_root = workspace / "files"
    files_root.mkdir(parents=True)
    profile = SimpleNamespace(
        agent_runtime=AgentRuntimeConfig(),
        config_for_role=lambda role: SimpleNamespace(model=role, provider="langchain", base_url=None),
    )
    runner = build_specialist_runner(
        workspace=workspace, llm_profile=profile, reporter=None,
        run_control=None, project_id="presentation", preferred_entrypoint="writing",
    ).runner
    backend = runner._make_backend(files_root=files_root, store=InMemoryStore())
    content = backend.read("/.easyslides/references/example.md")
    assert content.error is None
    assert "Library reference content" in content.file_data["content"]
    result = backend.execute('python "$CATMASTER_EASYSLIDES_ROOT/scripts/example.py"')
    assert result.exit_code == 0
    assert "Library reference content" in result.output
    assert (files_root / "output.txt").read_text() == "workspace output"


@pytest.mark.skipif(not RUNTIME_ROOT.is_dir(), reason="Install EasySlides for its export integration test")
def test_installed_export_retains_native_text_shapes_and_notes(tmp_path: Path) -> None:
    project = tmp_path / "talk"
    (project / "svg_output").mkdir(parents=True)
    (project / "notes").mkdir()
    (project / "svg_output/01_cover.svg").write_text('''<svg xmlns="http://www.w3.org/2000/svg" width="1280" height="720" viewBox="0 0 1280 720">
<rect width="1280" height="720" fill="#ffffff"/>
<text x="80" y="140" font-family="Arial" font-size="48" fill="#123456">Editable title</text>
<rect x="80" y="220" width="350" height="160" fill="#dcefff"/>
<text x="105" y="300" font-family="Arial" font-size="30">Editable body</text>
</svg>''')
    (project / "notes/01_cover.md").write_text("Speaker notes for the title slide.")
    output = tmp_path / "talk.pptx"
    result = subprocess.run(
        [sys.executable, str(RUNTIME_ROOT / "scripts/svg_to_pptx.py"), str(project),
         "--only", "native", "-t", "none", "-a", "none", "-o", str(output)],
        text=True, capture_output=True, timeout=300,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    deck = Presentation(output)
    slide = deck.slides[0]
    assert not any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes)
    title = next(shape for shape in slide.shapes if shape.has_text_frame and shape.text == "Editable title")
    assert "Speaker notes for the title slide." in slide.notes_slide.notes_text_frame.text
    title.text = "Title edited after export"
    deck.save(output)
    assert any(shape.has_text_frame and shape.text == "Title edited after export" for shape in Presentation(output).slides[0].shapes)
