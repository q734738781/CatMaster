from __future__ import annotations

import base64
from pathlib import Path
from types import SimpleNamespace

from langchain_core.messages import AIMessage, ToolMessage

import catmaster.runtime.document_reads as document_reads
from catmaster.runtime.checkpoint_serde import FileSafeCheckpointSerializer
from catmaster.runtime.deepagents_backend import CatMasterLocalShellBackend
from catmaster.runtime.document_reads import (
    BoundedDocumentReadMiddleware,
    native_document_input_decision,
    read_bounded_document,
)
from catmaster.runtime.multimodal_blocks import ModelMultimodalCapability


def test_native_backend_returns_office_files_as_provider_file_bytes(
    tmp_path: Path,
) -> None:
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pypdf import PdfWriter

    docx = Document()
    docx.add_paragraph("compact report")
    docx.save(tmp_path / "report.docx")

    workbook = Workbook()
    workbook.active.append(["name", "value"])
    workbook.active.append(["sample", 1])
    workbook.save(tmp_path / "results.xlsx")

    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(tmp_path / "slides.pptx")

    pdf = PdfWriter()
    pdf.add_blank_page(width=72, height=72)
    with (tmp_path / "paper.pdf").open("wb") as handle:
        pdf.write(handle)

    backend = CatMasterLocalShellBackend(root_dir=tmp_path, virtual_mode=True)
    for name in ("report.docx", "results.xlsx", "slides.pptx", "paper.pdf"):
        payload = (tmp_path / name).read_bytes()

        result = backend.read(f"/{name}")

        assert result.error is None
        assert result.file_data is not None
        assert result.file_data["encoding"] == "base64"
        assert base64.b64decode(result.file_data["content"]) == payload


def test_large_compressible_docx_is_not_returned_as_one_native_payload(
    tmp_path: Path,
) -> None:
    from docx import Document

    path = tmp_path / "large.docx"
    document = Document()
    document.add_paragraph("large document text " * 8_000)
    document.save(path)
    assert path.stat().st_size < 8 * 1024 * 1024
    assert native_document_input_decision(path).allowed is False

    result = CatMasterLocalShellBackend(
        root_dir=tmp_path,
        virtual_mode=True,
    ).read("/large.docx")

    assert result.file_data is None
    assert "bounded" in str(result.error)


def test_bounded_docx_read_uses_explicit_line_offsets_without_cursor_or_hash(
    tmp_path: Path,
) -> None:
    from docx import Document

    document = Document()
    for index in range(12):
        document.add_paragraph(f"paragraph {index}")
    document.save(tmp_path / "many.docx")

    first = read_bounded_document(
        tmp_path,
        file_path="/many.docx",
        offset=0,
        limit=3,
    )
    second = read_bounded_document(
        tmp_path,
        file_path="/many.docx",
        offset=first.next_offset,
        limit=3,
    )

    assert first.next_offset == 3
    assert "paragraph 0" in first.content
    assert "paragraph 3" not in first.content
    assert "paragraph 3" in second.content
    assert "cursor" not in first.content.lower()
    assert "hash" not in first.content.lower()
    assert "base64" not in first.content.lower()


def test_bounded_document_pages_reuse_incremental_parse_and_refresh_after_edit(
    tmp_path: Path,
    monkeypatch,
) -> None:
    path = tmp_path / "cached.docx"
    path.write_text("first version", encoding="utf-8")
    parse_calls: list[str] = []

    def fake_document_lines(source: Path):
        marker = source.read_text(encoding="utf-8")
        parse_calls.append(marker)
        for index in range(120):
            yield f"{marker} line {index}"

    monkeypatch.setattr(document_reads, "_iter_document_lines", fake_document_lines)

    assert native_document_input_decision(path).allowed is False
    first = read_bounded_document(
        tmp_path,
        file_path="/cached.docx",
        offset=0,
        limit=2,
    )
    second = read_bounded_document(
        tmp_path,
        file_path="/cached.docx",
        offset=first.next_offset,
        limit=2,
    )

    assert parse_calls == ["first version"]
    assert "first version line 2" in second.content

    path.write_text("second version is different", encoding="utf-8")
    refreshed = read_bounded_document(
        tmp_path,
        file_path="/cached.docx",
        offset=0,
        limit=2,
    )

    assert parse_calls == ["first version", "second version is different"]
    assert "second version is different line 0" in refreshed.content


def test_bounded_pdf_pptx_and_xlsx_reads_keep_remaining_content_reachable(
    tmp_path: Path,
) -> None:
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches
    from pypdf import PdfWriter

    pdf = PdfWriter()
    pdf.add_blank_page(width=72, height=72)
    pdf.add_blank_page(width=72, height=72)
    with (tmp_path / "pages.pdf").open("wb") as handle:
        pdf.write(handle)

    presentation = Presentation()
    for label in ("first slide", "second slide"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        box.text_frame.text = label
    presentation.save(tmp_path / "slides.pptx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["row", 1])
    sheet.append(["row", 2])
    workbook.save(tmp_path / "rows.xlsx")

    for file_path in ("/pages.pdf", "/slides.pptx", "/rows.xlsx"):
        first = read_bounded_document(
            tmp_path,
            file_path=file_path,
            offset=0,
            limit=1,
        )
        second = read_bounded_document(
            tmp_path,
            file_path=file_path,
            offset=first.next_offset,
            limit=1,
        )

        assert first.returned_lines == 1
        assert first.next_offset == 1
        assert second.returned_lines == 1
        assert second.content != first.content


def test_document_boundary_keeps_one_read_file_tool_and_intercepts_only_when_needed(
    tmp_path: Path,
) -> None:
    from docx import Document

    compact = Document()
    compact.add_paragraph("small")
    compact.save(tmp_path / "compact.docx")

    large = Document()
    large.add_paragraph("large document text " * 8_000)
    large.save(tmp_path / "large.docx")

    middleware = BoundedDocumentReadMiddleware(files_root=tmp_path)
    handler_result = ToolMessage(
        content="native",
        name="read_file",
        tool_call_id="native-call",
    )

    compact_request = SimpleNamespace(
        tool_call={
            "name": "read_file",
            "id": "compact-call",
            "args": {"file_path": "/compact.docx"},
        }
    )
    compact_result = middleware.wrap_tool_call(
        compact_request,
        lambda _request: handler_result,
    )
    assert compact_result is handler_result

    large_request = SimpleNamespace(
        tool_call={
            "name": "read_file",
            "id": "large-call",
            "args": {"file_path": "/large.docx"},
        }
    )
    large_result = middleware.wrap_tool_call(
        large_request,
        lambda _request: (_ for _ in ()).throw(AssertionError("native handler called")),
    )
    assert isinstance(large_result, ToolMessage)
    assert large_result.name == "read_file"
    assert large_result.status == "success"
    assert large_result.additional_kwargs["catmaster_bounded_document_read"] is True
    assert "Bounded DOCX text view" in str(large_result.content)


def test_native_backend_preserves_normal_text_pagination(tmp_path: Path) -> None:
    (tmp_path / "notes.txt").write_text("one\ntwo\nthree\n", encoding="utf-8")
    backend = CatMasterLocalShellBackend(root_dir=tmp_path, virtual_mode=True)

    result = backend.read("/notes.txt", offset=1, limit=1)

    assert result.error is None
    assert result.file_data == {"content": "two\n", "encoding": "utf-8"}
    assert result.next_offset == 2


def test_native_backend_treats_files_prefix_as_workspace_root_alias(
    tmp_path: Path,
) -> None:
    backend = CatMasterLocalShellBackend(root_dir=tmp_path, virtual_mode=True)

    written = backend.write(
        "files/reports/co2_fts_reproduction_plan.md",
        "real report\n",
    )
    reread = backend.read("/files/reports/co2_fts_reproduction_plan.md")

    assert written.error is None
    assert (tmp_path / "reports" / "co2_fts_reproduction_plan.md").read_text(
        encoding="utf-8"
    ) == "real report\n"
    assert not (tmp_path / "files").exists()
    assert reread.error is None
    assert reread.file_data == {"content": "real report\n", "encoding": "utf-8"}

    (tmp_path / "keep.txt").write_text("keep", encoding="utf-8")
    refused_root_alias = backend.delete("/files")
    assert refused_root_alias.error is not None
    assert (tmp_path / "keep.txt").is_file()


def test_bounded_document_read_accepts_user_visible_files_prefix(
    tmp_path: Path,
) -> None:
    from docx import Document

    reports = tmp_path / "reports"
    reports.mkdir()
    document = Document()
    document.add_paragraph("aliased report")
    document.save(reports / "report.docx")

    page = read_bounded_document(
        tmp_path,
        file_path="/files/reports/report.docx",
        offset=0,
        limit=5,
    )

    assert "aliased report" in page.content


def test_codex_oauth_profile_enables_native_document_blocks() -> None:
    capability = ModelMultimodalCapability.from_llm_config(
        type("Config", (), {"provider": "codex_oauth", "provider_options": {}})()
    )

    assert capability.pdfs is True
    assert capability.documents is True


def test_checkpoint_serializer_replaces_native_file_bytes_with_reopen_hint() -> None:
    original = ToolMessage(
        content_blocks=[
            {
                "type": "file",
                "base64": "UEsDBHJlcG9ydA==",
                "mime_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            }
        ],
        additional_kwargs={
            "read_file_path": "/reports/report.docx",
            "read_file_media_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        },
        tool_call_id="call-docx",
        name="read_file",
    )
    serializer = FileSafeCheckpointSerializer()

    type_name, payload = serializer.dumps_typed(
        {"channel_values": {"messages": [original]}}
    )
    restored = serializer.loads_typed((type_name, payload))
    message = restored["channel_values"]["messages"][0]

    assert b"UEsDBHJlcG9ydA" not in payload
    assert isinstance(message, ToolMessage)
    assert message.tool_call_id == "call-docx"
    assert message.additional_kwargs["catmaster_inline_file_payload_removed"] is True
    assert "native `read_file`" in str(message.content)
    assert "/reports/report.docx" in str(message.content)


def test_checkpoint_serializer_preserves_all_inline_images() -> None:
    messages = []
    for index in range(5):
        messages.extend(
            [
                ToolMessage(
                    content_blocks=[
                        {
                            "type": "image",
                            "base64": f"payload-{index}",
                            "mime_type": "image/png",
                        }
                    ],
                    additional_kwargs={"read_file_path": f"/figure-{index}.png"},
                    tool_call_id=f"call-image-{index}",
                    name="read_file",
                ),
                AIMessage(content=f"inspected image {index}"),
            ]
        )
    serializer = FileSafeCheckpointSerializer()

    restored = serializer.loads_typed(serializer.dumps_typed(messages))
    image_messages = [
        message for message in restored if isinstance(message, ToolMessage)
    ]

    assert [message.content[0]["base64"] for message in image_messages] == [
        f"payload-{index}" for index in range(5)
    ]
