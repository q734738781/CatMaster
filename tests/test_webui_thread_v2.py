from __future__ import annotations

import base64
import asyncio
import json
import multiprocessing
import time
import uuid
from contextlib import asynccontextmanager
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pytest
from fastapi.testclient import TestClient
from langchain_core.messages import AIMessage, AIMessageChunk, ToolMessage
from langgraph.types import Command

from catmaster.research.knowledge_graph.store import ResearchGraphStore
from catmaster.runtime.run_context import RunContext
from catmaster.tools.base import ensure_project_space_layout, system_root
from catmaster.runtime.observability_store import OBSERVABILITY_DB_NAME, ObservabilityStore
from catmaster.webui.agent_loop import ThreadAgentLoopService
from catmaster.webui import server
from catmaster.webui.artifact_registry import ArtifactRegistry, infer_renderer
from catmaster.webui.projections import project_event
from catmaster.webui.projections.tools import project_tool_part
from catmaster.webui.server import create_app as _create_app
from catmaster.webui.thread_events import ThreadEventBroker
from catmaster.webui.thread_models import (
    ArtifactPart,
    MessagePart,
    ThreadCheckpointContinueRequest,
    ThreadMessage,
    ThreadStopRequest,
    ThreadSubmitRequest,
)
from catmaster.webui.thread_store import ThreadStore, new_id
from catmaster.specialists.runtime import SpecialistUsageCallbackHandler


class _TestExecutionHost:
    """Route tests exercise native graph projection without a background executor.

    Durable scheduling is tested separately against real DBOS.
    """
    def __init__(self):
        self.rows = {}
        self.calls = []
        self.factory = None

    async def start(self):
        pass

    async def close(self):
        pass

    async def runs(self, workspace, thread_id, *, active=False, **kwargs):
        return [] if active else list(self.rows.values())

    async def run(self, run_id):
        return self.rows.get(run_id)

    async def enqueue(self, packet):
        from deepagents import create_deep_agent
        from langgraph.checkpoint.sqlite.aio import AsyncSqliteSaver
        from langchain_core.language_models.fake_chat_models import FakeMessagesListChatModel
        class Model(FakeMessagesListChatModel):
            def bind_tools(self, tools, **kwargs):
                return self
        service = self.factory(Path(packet['workspace']), packet['workspace_id'])
        host = self
        @asynccontextmanager
        async def graph_factory(_service, _packet):
            async with AsyncSqliteSaver.from_conn_string(str(service.workspace / 'metadata/deepagent_threads.sqlite')) as saver:
                graph = create_deep_agent(Model(responses=[AIMessage(content='done')]), checkpointer=saver)
                yield graph
                state = await graph.aget_state({'configurable': {'thread_id': service.store.get_thread(packet['thread_id']).deepagent_thread_id}})
                host.calls.append({'input': {'messages': [m.model_dump(mode='json') for m in state.values['messages'] if m.type == 'human']}})
        service.graph_factory = graph_factory
        result = await service.execute_turn(packet)
        self.rows[packet['run_id']] = SimpleNamespace(workflow_id=packet['run_id'], status='SUCCESS', output=result)
        return packet['run_id']


def create_app(**kwargs: Any):
    fake = kwargs.setdefault('execution_host', _TestExecutionHost())
    app = _create_app(**kwargs)
    app.state.test_execution = fake
    return app


def _register_artifact_process(
    workspace: str,
    path: str,
    thread_id: str,
    barrier: Any,
    results: Any,
) -> None:
    try:
        barrier.wait(timeout=10)
        record = ArtifactRegistry(
            workspace=Path(workspace),
            workspace_id="default",
        ).register_path(path, thread_id=thread_id)
        results.put({"artifact_id": record.artifact_id})
    except Exception as exc:  # pragma: no cover - surfaced in the parent
        results.put({"error": f"{type(exc).__name__}: {exc}"})


def _wait_for_thread_event_process(
    workspace: str,
    thread_id: str,
    ready: Any,
    results: Any,
) -> None:
    try:
        broker = ThreadEventBroker(workspace=Path(workspace))
        ready.set()
        events, cursor = broker.wait_for_events(
            thread_id,
            last_seq=0,
            timeout_s=5,
        )
        results.put(
            {
                "events": [event.event for event in events],
                "cursor": cursor,
            }
        )
    except Exception as exc:  # pragma: no cover - surfaced in the parent
        results.put({"error": f"{type(exc).__name__}: {exc}"})


def _workspace(tmp_path: Path) -> Path:
    workspace = tmp_path / "default"
    ensure_project_space_layout(workspace, create=True)
    return workspace


def test_thread_request_models_keep_model_config_as_an_api_alias() -> None:
    default_request = ThreadSubmitRequest(text="hello")
    selected_request = ThreadSubmitRequest(text="hello", model_config="configs/custom.yaml")

    assert default_request.llm_config == ""
    assert selected_request.llm_config == "configs/custom.yaml"
    assert ThreadSubmitRequest(text="hello", llm_config="configs/by-name.yaml").llm_config == "configs/by-name.yaml"
    assert "model_config" in ThreadSubmitRequest.model_json_schema()["properties"]
    assert "llm_config" not in ThreadSubmitRequest.model_json_schema()["properties"]
    assert ThreadStopRequest().run_id == ""
    assert ThreadStopRequest().action == "interrupt"
    assert ThreadStopRequest().reason == ""
    assert ThreadCheckpointContinueRequest(message_id="msg_failed").message_id == "msg_failed"


def test_thread_workspace_image_route_serves_only_images_from_the_thread_workspace(
    tmp_path: Path,
) -> None:
    workspace = _workspace(tmp_path)
    store = ThreadStore(workspace=workspace, workspace_id="default")
    thread = store.create_thread(title="Inline image")
    image_path = workspace / "files" / "figures" / "curve.png"
    image_path.parent.mkdir(parents=True, exist_ok=True)
    image_bytes = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    )
    image_path.write_bytes(image_bytes)
    (workspace / "files" / "figures" / "notes.txt").write_text("not an image", encoding="utf-8")

    client = TestClient(create_app(project_space_root=str(tmp_path), no_login=True))
    response = client.get(
        f"/api/threads/{thread.thread_id}/files/image",
        params={"path": "files/figures/curve.png"},
    )

    assert response.status_code == 200
    assert response.content == image_bytes
    assert response.headers["content-type"] == "image/png"
    assert response.headers["cache-control"] == "private, max-age=0, must-revalidate"
    assert client.get(
        f"/api/threads/{thread.thread_id}/files/image",
        params={"path": "files/figures/notes.txt"},
    ).status_code == 415
    assert client.get(
        f"/api/threads/{thread.thread_id}/files/image",
        params={"path": "metadata/workspace.sqlite"},
    ).status_code == 404


def test_thread_list_warms_workspace_lookup_for_thread_routes(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workspace = _workspace(tmp_path)
    thread = ThreadStore(workspace=workspace, workspace_id="default").create_thread(
        title="Cached workspace routing"
    )
    client = TestClient(create_app(project_space_root=str(tmp_path), no_login=True))

    listed = client.get("/api/workspaces/default/threads")
    assert listed.status_code == 200
    assert thread.thread_id in {
        item["thread_id"] for item in listed.json()["threads"]
    }

    original_rglob = Path.rglob
    reverse_scans: list[str] = []

    def tracking_rglob(path: Path, pattern: str):
        if pattern == f"metadata/threads/{thread.thread_id}/thread.json":
            reverse_scans.append(pattern)
        return original_rglob(path, pattern)

    monkeypatch.setattr(Path, "rglob", tracking_rglob)
    assert client.get(f"/api/threads/{thread.thread_id}/messages").status_code == 200
    assert client.get(f"/api/threads/{thread.thread_id}/artifacts").status_code == 200
    assert reverse_scans == []


def test_checkpoint_continue_route_forwards_failed_message_identity(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workspace = _workspace(tmp_path)
    client = TestClient(create_app(project_space_root=str(tmp_path), no_login=True))
    created = client.post(
        "/api/workspaces/default/threads",
        json={"title": "Checkpoint continuation"},
    )
    thread_id = created.json()["thread"]["thread_id"]
    captured: dict[str, str] = {}

    async def fake_continue(self, *, thread_id: str, payload: Any):
        captured["thread_id"] = thread_id
        captured["message_id"] = payload.message_id
        message = ThreadMessage(
            id="msg_checkpoint_continuation",
            thread_id=thread_id,
            role="assistant",
            status="streaming",
            parts=[
                MessagePart(
                    id="part_checkpoint_continuation",
                    type="text",
                    status="streaming",
                )
            ],
        )
        self.store.append_message(message)
        thread = self.store.update_thread(thread_id, status="running")
        return {
            "accepted": True,
            "assistant_message": message,
            "thread": thread,
        }

    monkeypatch.setattr(
        ThreadAgentLoopService,
        "continue_from_checkpoint",
        fake_continue,
    )
    response = client.post(
        f"/api/threads/{thread_id}/continue-from-checkpoint",
        json={"message_id": "msg_failed_route"},
    )

    assert response.status_code == 200
    assert captured == {
        "thread_id": thread_id,
        "message_id": "msg_failed_route",
    }
    assert response.json()["assistant_message"]["id"] == "msg_checkpoint_continuation"
    assert client.post(
        f"/api/threads/{thread_id}/continue-from-checkpoint",
        json={},
    ).status_code == 422


def test_thread_store_persists_messages_and_events_replay(tmp_path: Path) -> None:
    workspace = _workspace(tmp_path)
    store = ThreadStore(workspace=workspace, workspace_id="default")
    thread = store.create_thread(title="hello")
    message = ThreadMessage(
        id=new_id("msg"),
        thread_id=thread.thread_id,
        role="user",
        status="completed",
        parts=[MessagePart(id=new_id("part_text"), type="text", text="hi", status="completed")],
    )
    store.append_message(message)

    assert store.get_thread(thread.thread_id).title == "hello"
    assert store.list_messages(thread.thread_id)[0].parts[0].text == "hi"

    broker = ThreadEventBroker(workspace=workspace)
    first = broker.emit(thread.thread_id, "message.created", message_id=message.id, data={"message_id": message.id})
    second = broker.emit(thread.thread_id, "message.completed", message_id=message.id)

    assert first.seq == 1
    assert second.seq == 2
    replay = broker.replay(thread.thread_id, last_seq=1)
    assert [event.event for event in replay] == ["message.completed"]
    assert not broker.events_path(thread.thread_id).exists()


def test_thread_event_broker_persists_stream_events_to_observability_store(tmp_path: Path) -> None:
    workspace = _workspace(tmp_path)
    store = ThreadStore(workspace=workspace, workspace_id="default")
    thread = store.create_thread(title="observe")
    run_dir = system_root(workspace) / "runs" / "run_observe"
    RunContext.create(
        workspace=workspace,
        project_id="default",
        run_id="run_observe",
        model_name="test-model",
    )

    broker = ThreadEventBroker(workspace=workspace)
    broker.emit(thread.thread_id, "reasoning.delta", message_id="msg_1", data={"run_id": "run_observe", "part_id": "part_reasoning"})

    assert (run_dir / OBSERVABILITY_DB_NAME).exists()
    rows = ObservabilityStore(run_dir).read_thread_events_page(thread.thread_id)
    assert [row["name"] for row in rows] == ["reasoning.delta"]
    assert rows[0]["channel"] == "thread"
    assert rows[0]["message_id"] == "msg_1"

    restarted = ThreadEventBroker(workspace=workspace)
    replay = restarted.replay(thread.thread_id)
    assert [event.event for event in replay] == ["reasoning.delta"]


def test_thread_event_broker_does_not_claim_native_run_directory(tmp_path: Path) -> None:
    workspace = _workspace(tmp_path)
    thread = ThreadStore(
        workspace=workspace,
        workspace_id="default",
    ).create_thread(title="native run allocation")
    run_dir = system_root(workspace) / "runs" / "run_native_pending"

    event = ThreadEventBroker(workspace=workspace).emit(
        thread.thread_id,
        "message.created",
        message_id="msg_pending",
        data={"run_id": "run_native_pending"},
    )

    assert not run_dir.exists()
    assert event.event == "message.created"
    replay = ThreadEventBroker(workspace=workspace).replay(thread.thread_id)
    assert [row.event for row in replay] == ["message.created"]


def test_thread_event_broker_wakes_a_different_process(tmp_path: Path) -> None:
    workspace = _workspace(tmp_path)
    thread_id = ThreadStore(
        workspace=workspace,
        workspace_id="default",
    ).create_thread(title="cross-process").thread_id
    context = multiprocessing.get_context("spawn")
    ready = context.Event()
    results = context.Queue()
    process = context.Process(
        target=_wait_for_thread_event_process,
        args=(str(workspace), thread_id, ready, results),
    )
    process.start()
    assert ready.wait(timeout=10)

    emitted = ThreadEventBroker(workspace=workspace).emit(
        thread_id,
        "message.created",
        message_id="msg_cross_process",
    )
    outcome = results.get(timeout=10)
    process.join(timeout=10)

    assert process.exitcode == 0
    assert outcome == {
        "events": ["message.created"],
        "cursor": emitted.seq,
    }


def test_thread_stream_route_replays_observability_events_after_reconnect(tmp_path: Path) -> None:
    workspace = _workspace(tmp_path)
    app = create_app(project_space_root=str(tmp_path), no_login=True)
    client = TestClient(app)
    created = client.post("/api/workspaces/default/threads", json={"title": "SSE replay"})
    assert created.status_code == 200
    thread_id = created.json()["thread"]["thread_id"]

    broker = ThreadEventBroker(workspace=workspace)
    broker.emit(thread_id, "reasoning.delta", message_id="msg_1", data={"run_id": "run_sse", "part_id": "part_reasoning"})
    broker.emit(thread_id, "message.delta", message_id="msg_1", data={"run_id": "run_sse", "part_id": "part_text"})

    response = client.get(f"/api/threads/{thread_id}/stream", params={"last_seq": "1", "once": "true"})

    assert response.status_code == 200
    assert "event: message.delta" in response.text

    snapshot = client.get(f"/api/threads/{thread_id}/messages").json()
    cursor = snapshot["stream_cursor"]
    assert cursor == broker.latest_seq(thread_id)
    next_event = broker.emit(thread_id, "message.delta", message_id="msg_1",
        data={"part_id": "part_text", "delta": "409", "text_offset": 6})
    replay = client.get(f"/api/threads/{thread_id}/stream",
        params={"last_seq": "0", "once": "true"}, headers={"Last-Event-ID": str(cursor)})
    assert f"id: {next_event.seq}" in replay.text
    assert '"text_offset":6' in replay.text
    assert "reasoning.delta" not in replay.text


def test_thread_stream_deduplicates_the_same_failure_across_reconnects(
    tmp_path: Path,
) -> None:
    workspace = _workspace(tmp_path)
    app = create_app(project_space_root=str(tmp_path), no_login=True)
    client = TestClient(app)
    thread_id = client.post(
        "/api/workspaces/default/threads",
        json={"title": "SSE failure reconnect"},
    ).json()["thread"]["thread_id"]
    broker = ThreadEventBroker(workspace=workspace)
    first = broker.emit(
        thread_id,
        "message.failed",
        message_id="msg_failure_domain",
        data={"error": "Remote calculation failed."},
    )

    initial = client.get(
        f"/api/threads/{thread_id}/stream",
        params={"last_seq": str(first.seq - 1), "once": "true"},
    )
    assert initial.text.count("event: run.failed") == 1

    duplicate = broker.emit(
        thread_id,
        "error",
        message_id="msg_failure_domain",
        data={"error": "Remote calculation failed."},
    )
    reconnected = client.get(
        f"/api/threads/{thread_id}/stream",
        params={"last_seq": str(first.seq), "once": "true"},
        headers={"Last-Event-ID": str(first.seq)},
    )
    assert duplicate.seq > first.seq
    assert reconnected.text.count("event: run.failed") == 0

    broker.emit(
        thread_id,
        "thread.status",
        message_id="msg_failure_domain",
        status="running",
        data={"status": "running"},
    )
    next_failure = broker.emit(
        thread_id,
        "message.failed",
        message_id="msg_failure_domain",
        data={"error": "A later retry failed."},
    )
    retried = client.get(
        f"/api/threads/{thread_id}/stream",
        params={"last_seq": str(duplicate.seq), "once": "true"},
    )
    assert next_failure.seq > duplicate.seq
    assert retried.text.count("event: run.failed") == 1


def test_thread_store_migrates_legacy_chat_sessions(tmp_path: Path) -> None:
    workspace = _workspace(tmp_path)
    legacy_session = system_root(workspace) / "chat_sessions" / "session_a"
    legacy_session.mkdir(parents=True)
    (legacy_session / "session.json").write_text(json.dumps({"title": "Legacy chat"}), encoding="utf-8")
    rows = [
        {"message_id": "msg_user", "role": "user", "content": "Run O2.", "created_at": 10.0},
        {
            "message_id": "msg_result",
            "role": "assistant",
            "kind": "run_result",
            "source_run_id": "run_o2",
            "content": "O2 done.",
            "created_at": 11.0,
        },
    ]
    (legacy_session / "messages.jsonl").write_text(
        "\n".join(json.dumps(row) for row in rows) + "\n",
        encoding="utf-8",
    )

    store = ThreadStore(workspace=workspace, workspace_id="default")
    thread = store.get_thread("thread_session_a")
    messages = store.list_messages(thread.thread_id)

    assert thread.title == "Legacy chat"
    assert thread.meta["legacy_chat_session_id"] == "session_a"
    assert [message.role for message in messages] == ["user", "assistant"]
    assert messages[0].parts[0].text == "Run O2."
    assert messages[1].parts[0].text == "O2 done."
    assert messages[1].meta["legacy_kind"] == "run_result"
    assert messages[1].meta["run_id"] == "run_o2"


def test_artifact_registry_renderer_and_path_safety(tmp_path: Path) -> None:
    workspace = _workspace(tmp_path)
    (workspace / "files" / "table.csv").write_text("a,b\n1,2\n", encoding="utf-8")
    registry = ArtifactRegistry(workspace=workspace, workspace_id="default")

    record = registry.register_path("table.csv", thread_id="thread_x", message_id="msg_x")

    assert record.path == "files/table.csv"
    assert record.renderer == "csv"
    assert infer_renderer("POSCAR") == "structure"
    assert infer_renderer("figure.png", "image/png") == "image"

    try:
        registry.register_path("../secret.txt")
    except ValueError as exc:
        assert "invalid" in str(exc) or "escapes" in str(exc)
    else:
        raise AssertionError("unsafe artifact path was accepted")


def test_artifact_registry_reuses_one_thread_artifact_for_the_same_path(
    tmp_path: Path,
) -> None:
    workspace = _workspace(tmp_path)
    (workspace / "files" / "table.csv").write_text("a,b\n1,2\n", encoding="utf-8")
    registry = ArtifactRegistry(workspace=workspace, workspace_id="default")

    first = registry.register_path(
        "table.csv",
        thread_id="thread_x",
        message_id="msg_one",
        tool_call_id="call_one",
    )
    second = registry.register_path(
        "table.csv",
        thread_id="thread_x",
        message_id="msg_two",
        tool_call_id="call_two",
    )

    assert second.artifact_id == first.artifact_id
    assert uuid.UUID(first.artifact_id.removeprefix("art_"))
    assert [record.artifact_id for record in registry.list_artifacts(thread_id="thread_x")] == [
        first.artifact_id
    ]


def test_artifact_registry_imports_legacy_index_once_without_rewriting_it(tmp_path: Path) -> None:
    workspace = _workspace(tmp_path)
    output = workspace / "files" / "legacy-index.txt"
    output.write_text("legacy", encoding="utf-8")
    legacy_root = system_root(workspace) / "artifacts"
    legacy_root.mkdir(parents=True)
    index_path = legacy_root / "index.jsonl"
    original = (
        json.dumps(
            {
                "artifact_id": "art_legacy_index",
                "thread_id": "thread_legacy",
                "workspace_id": "default",
                "path": "files/legacy-index.txt",
                "title": "Legacy index",
            }
        )
        + "\n"
    )
    index_path.write_text(original, encoding="utf-8")

    first = ArtifactRegistry(workspace=workspace, workspace_id="default")
    second = ArtifactRegistry(workspace=workspace, workspace_id="default")

    assert first.get("art_legacy_index") is not None
    assert [item.artifact_id for item in second.list_artifacts()] == [
        "art_legacy_index"
    ]
    assert index_path.read_text(encoding="utf-8") == original


def test_artifact_registry_concurrent_processes_do_not_lose_updates(tmp_path: Path) -> None:
    workspace = _workspace(tmp_path)
    paths = ["first.txt", "second.txt"]
    for path in paths:
        (workspace / "files" / path).write_text(path, encoding="utf-8")
    context = multiprocessing.get_context("spawn")
    barrier = context.Barrier(len(paths) + 1)
    results = context.Queue()
    processes = [
        context.Process(
            target=_register_artifact_process,
            args=(str(workspace), path, f"thread_{index}", barrier, results),
        )
        for index, path in enumerate(paths)
    ]
    for process in processes:
        process.start()
    barrier.wait(timeout=10)
    outcomes = [results.get(timeout=10) for _ in processes]
    for process in processes:
        process.join(timeout=10)

    assert [process.exitcode for process in processes] == [0, 0]
    assert all("error" not in outcome for outcome in outcomes)
    records = ArtifactRegistry(
        workspace=workspace,
        workspace_id="default",
    ).list_artifacts()
    assert {record.path for record in records} == {
        "files/first.txt",
        "files/second.txt",
    }
    assert not (system_root(workspace) / "artifacts" / "index.jsonl").exists()


def test_artifact_registry_renderer_mapping_covers_domain_artifacts(tmp_path: Path) -> None:
    assert infer_renderer("POSCAR") == "structure"
    assert infer_renderer("trajectory.traj") == "structure"
    assert infer_renderer("figure.svg", "image/svg+xml") == "image"
    assert infer_renderer("table.tsv") == "csv"
    assert infer_renderer("report.rst") == "markdown"
    assert infer_renderer("paper.pdf", "application/pdf") == "pdf"
    assert infer_renderer("stdout.log") == "text"
    assert infer_renderer("bundle.zip") == "archive"


def test_artifact_registry_skips_missing_run_state_artifacts(tmp_path: Path) -> None:
    workspace = _workspace(tmp_path)
    (workspace / "files" / "notes").mkdir(parents=True)
    (workspace / "files" / "notes" / "summary.json").write_text("{}", encoding="utf-8")
    registry = ArtifactRegistry(workspace=workspace, workspace_id="default")

    records = registry.register_from_run_state(
        {
            "thread_id": "thread_x",
            "run_id": "run_x",
            "artifacts": [
                {"path": "fmax=0.02 eV/Å", "description": "not a file"},
                {"path": "notes/summary.json", "description": "real output"},
            ],
        },
        thread_id="thread_x",
        run_id="run_x",
    )

    assert [record.path for record in records] == ["files/notes/summary.json"]


def test_artifact_registry_rejects_missing_tool_artifact_paths(tmp_path: Path) -> None:
    workspace = _workspace(tmp_path)
    registry = ArtifactRegistry(workspace=workspace, workspace_id="default")

    with pytest.raises(ValueError, match="existing workspace file"):
        registry.register_path(
            "task_config.fmax",
            thread_id="thread_x",
            tool_call_id="call_task_spec",
            meta={"source": "tool_artifact"},
        )

    assert registry.list_artifacts(thread_id="thread_x") == []


def test_artifact_registry_hides_index_records_after_file_disappears(tmp_path: Path) -> None:
    workspace = _workspace(tmp_path)
    output = workspace / "files" / "temporary.csv"
    output.write_text("x\n1\n", encoding="utf-8")
    registry = ArtifactRegistry(workspace=workspace, workspace_id="default")
    record = registry.register_path("temporary.csv", thread_id="thread_x")

    output.unlink()

    assert registry.get(record.artifact_id) is None
    assert registry.list_artifacts(thread_id="thread_x") == []


def test_artifact_registry_migrates_legacy_run_artifacts(tmp_path: Path) -> None:
    workspace = _workspace(tmp_path)
    (workspace / "files" / "legacy.csv").write_text("x\n1\n", encoding="utf-8")
    run_dir = system_root(workspace) / "runs" / "run_legacy"
    run_dir.mkdir(parents=True)
    run_state_text = json.dumps(
        {
            "status": "done",
            "thread_id": "thread_legacy",
            "artifacts": [{"path": "legacy.csv", "summary": "legacy table"}],
        },
        indent=2,
    )
    (run_dir / "run_state.json").write_text(run_state_text, encoding="utf-8")
    checkpoint_path = system_root(workspace) / "deepagent_threads.sqlite"
    checkpoint_path.write_bytes(b"legacy checkpoint")
    registry = ArtifactRegistry(workspace=workspace, workspace_id="default")

    records = registry.list_artifacts(thread_id="thread_legacy")

    assert len(records) == 1
    assert records[0].path == "files/legacy.csv"
    assert records[0].renderer == "csv"
    assert records[0].run_id == "run_legacy"
    assert records[0].meta["source"] == "run_state"
    assert (run_dir / "run_state.json").read_text(encoding="utf-8") == run_state_text
    assert checkpoint_path.read_bytes() == b"legacy checkpoint"


def test_server_thread_routes_and_artifact_preview(tmp_path: Path) -> None:
    workspace = _workspace(tmp_path)
    (workspace / "files" / "note.md").write_text("# Result\n\nok\n", encoding="utf-8")
    app = create_app(project_space_root=str(tmp_path), no_login=True)
    client = TestClient(app)

    created = client.post("/api/workspaces/default/threads", json={"title": "T"})
    assert created.status_code == 200
    thread_id = created.json()["thread"]["thread_id"]

    listed = client.get("/api/workspaces/default/threads")
    assert listed.status_code == 200
    assert listed.json()["threads"][0]["thread_id"] == thread_id

    assert client.get(f"/api/threads/{thread_id}").status_code == 200
    empty_page = client.get(f"/api/threads/{thread_id}/messages").json()
    assert empty_page["messages"] == []
    assert empty_page["page"]["shown_count"] == 0
    assert empty_page["page"]["total_count"] == 0
    assert empty_page["page"]["total_unknown"] is False
    assert empty_page["page"]["truncated"] is False
    assert empty_page["page"]["next_cursor"] == ""

    registry = ArtifactRegistry(workspace=workspace, workspace_id="default")
    artifact = registry.register_path("note.md", thread_id=thread_id, message_id="msg_x")
    preview = client.get(f"/api/artifacts/{artifact.artifact_id}/preview")
    assert preview.status_code == 200
    assert preview.json()["kind"] == "markdown"
    assert "Result" in preview.json()["preview_text"]

    malformed = client.post(
        f"/api/threads/{thread_id}/resume",
        json={"actions": [{"action_id": "approval-1", "decision": "deny"}]},
    )
    assert malformed.status_code == 422


def test_server_hides_historical_artifact_parts_without_files(tmp_path: Path) -> None:
    workspace = _workspace(tmp_path)
    (workspace / "files" / "real.csv").write_text("x\n1\n", encoding="utf-8")
    app = create_app(project_space_root=str(tmp_path), no_login=True)
    client = TestClient(app)
    thread_id = client.post("/api/workspaces/default/threads", json={"title": "T"}).json()["thread"]["thread_id"]
    registry = ArtifactRegistry(workspace=workspace, workspace_id="default")
    real = registry.register_path("real.csv", thread_id=thread_id, message_id="msg_artifacts")
    store = ThreadStore(workspace=workspace, workspace_id="default")
    store.append_message(
        ThreadMessage(
            id="msg_artifacts",
            thread_id=thread_id,
            role="assistant",
            status="completed",
            parts=[
                ArtifactPart(
                    id="part_real",
                    artifact_id=real.artifact_id,
                    path=real.path,
                ),
                ArtifactPart(
                    id="part_missing",
                    artifact_id="art_task_config_fmax",
                    path="files/task_config.fmax",
                ),
            ],
        )
    )

    response = client.get(f"/api/threads/{thread_id}/messages")

    assert response.status_code == 200
    artifact_parts = [
        part
        for part in response.json()["messages"][0]["parts"]
        if part["type"] == "artifact"
    ]
    assert [part["artifact_id"] for part in artifact_parts] == [real.artifact_id]


def test_thread_permission_mode_create_and_patch(tmp_path: Path) -> None:
    _workspace(tmp_path)
    app = create_app(project_space_root=str(tmp_path), no_login=True)
    client = TestClient(app)

    default_created = client.post("/api/workspaces/default/threads", json={"title": "default"})
    assert default_created.status_code == 200
    default_thread = default_created.json()["thread"]
    assert default_thread["permission_mode"] == "auto"
    assert "meta" not in default_thread
    assert server._thread_permission_mode(SimpleNamespace(meta={})) == "auto"

    created = client.post("/api/workspaces/default/threads", json={"title": "auto", "permission_mode": "auto-approve"})
    assert created.status_code == 200
    thread = created.json()["thread"]
    assert thread["permission_mode"] == "auto"

    patched = client.patch(f"/api/threads/{thread['thread_id']}", json={"permission_mode": "review"})
    assert patched.status_code == 200
    thread = patched.json()["thread"]
    assert thread["permission_mode"] == "hitl"

    invalid = client.patch(f"/api/threads/{thread['thread_id']}", json={"permission_mode": "bad"})
    assert invalid.status_code == 400


def test_thread_title_patch_persists_user_rename(tmp_path: Path) -> None:
    _workspace(tmp_path)
    app = create_app(project_space_root=str(tmp_path), no_login=True)
    client = TestClient(app)

    created = client.post("/api/workspaces/default/threads", json={"title": "Initial title"})
    assert created.status_code == 200
    thread_id = created.json()["thread"]["thread_id"]

    renamed = client.patch(
        f"/api/threads/{thread_id}",
        json={"title": "ZrC potential comparison"},
    )
    assert renamed.status_code == 200
    assert renamed.json()["thread"]["title"] == "ZrC potential comparison"

    reopened = TestClient(create_app(project_space_root=str(tmp_path), no_login=True))
    fetched = reopened.get(f"/api/threads/{thread_id}")
    assert fetched.status_code == 200
    assert fetched.json()["thread"]["title"] == "ZrC potential comparison"


def test_thread_entrypoint_api_preserves_all_lanes_and_validates_values(tmp_path: Path) -> None:
    _workspace(tmp_path)
    app = create_app(project_space_root=str(tmp_path), no_login=True)
    client = TestClient(app)

    entrypoints = client.get("/api/entrypoints")
    assert entrypoints.status_code == 200
    ids = [item["id"] for item in entrypoints.json()["entrypoints"]]
    assert ids == [
        "research",
        "persistent_research",
        "experiment",
        "writing",
        "peer_review",
        "literature_review",
    ]

    created = client.post("/api/workspaces/default/threads", json={"title": "Write", "entrypoint": "writing"})
    assert created.status_code == 200
    thread = created.json()["thread"]
    assert thread["entrypoint"] == "writing"

    patched = client.patch(f"/api/threads/{thread['thread_id']}", json={"entrypoint": "peer-review"})
    assert patched.status_code == 200
    assert patched.json()["thread"]["entrypoint"] == "peer_review"

    alias = client.patch(f"/api/threads/{thread['thread_id']}", json={"entrypoint": "literature"})
    assert alias.status_code == 200
    assert alias.json()["thread"]["entrypoint"] == "literature_review"

    invalid = client.patch(f"/api/threads/{thread['thread_id']}", json={"entrypoint": "unknown"})
    assert invalid.status_code == 400


def test_submit_creates_user_and_running_assistant_message(tmp_path: Path) -> None:
    _workspace(tmp_path)
    app = create_app(project_space_root=str(tmp_path), no_login=True)
    client = TestClient(app)
    thread_id = client.post("/api/workspaces/default/threads", json={}).json()["thread"]["thread_id"]

    submitted = client.post(f"/api/threads/{thread_id}/submit", json={"text": "hello", "permission_mode": "auto"})

    assert submitted.status_code == 200
    payload = submitted.json()
    assert payload["message"]["role"] == "user"
    assert payload["assistant_message"]["role"] == "assistant"
    assert "meta" not in payload["assistant_message"]
    assert payload["thread"]["permission_mode"] == "auto"


@pytest.mark.parametrize("entrypoint", ["persistent_research", "writing", "experiment"])
def test_submit_without_entrypoint_retains_the_selected_lane(tmp_path: Path, entrypoint: str) -> None:
    _workspace(tmp_path)
    client = TestClient(create_app(project_space_root=str(tmp_path), no_login=True))
    thread = client.post("/api/workspaces/default/threads", json={"entrypoint": entrypoint}).json()["thread"]
    response = client.post(f"/api/threads/{thread['thread_id']}/submit", json={"text": "Continue this task."})
    assert response.status_code == 200
    assert response.json()["thread"]["entrypoint"] == entrypoint
    if entrypoint == "persistent_research":
        assert response.json()["thread"]["thread_role"] == "research_root"
        assert response.json()["thread"]["active_research_graph_id"]


def test_submit_image_attachment_registers_artifact_without_persisting_data_url(tmp_path: Path) -> None:
    workspace = _workspace(tmp_path)
    app = create_app(project_space_root=str(tmp_path), no_login=True)
    client = TestClient(app)
    thread_id = client.post("/api/workspaces/default/threads", json={}).json()["thread"]["thread_id"]
    image_data = base64.b64encode(b"fake-png").decode("ascii")
    submitted = client.post(
        f"/api/threads/{thread_id}/submit",
        json={
            "text": "inspect this image",
            "attachments": [
                {
                    "type": "image",
                    "filename": "figure.png",
                    "mime_type": "image/png",
                    "data": f"data:image/png;base64,{image_data}",
                }
            ],
        },
    )

    assert submitted.status_code == 200
    user_message = submitted.json()["message"]
    artifact_parts = [part for part in user_message["parts"] if part["type"] == "artifact"]
    assert len(artifact_parts) == 1
    assert artifact_parts[0]["path"].startswith("files/attachments/")
    assert "data:image/png" not in json.dumps(user_message)
    content = app.state.test_execution.calls[-1]["input"]["messages"][0]["content"]
    assert "data:image/png" not in str(content)
    assert isinstance(content, list)
    assert content[0]["type"] == "text"
    assert "figure.png" in content[0]["text"]
    assert len(content) == 2
    assert content[1] == {
        "type": "image",
        "base64": image_data,
        "mime_type": "image/png",
    }
    assert (workspace / artifact_parts[0]["path"]).exists()
    multimodal_events = [event for event in ThreadEventBroker(workspace=workspace).replay(thread_id) if event.event == "multimodal.prepared"]
    assert multimodal_events
    event_data = multimodal_events[-1].data
    assert event_data["attachment_count"] == 1
    assert event_data["attachments"][0]["sent_to_model"] is True
    assert event_data["attachments"][0]["sent_as"] == "image"
    assert event_data["attachments"][0]["warnings"] == []
    assert "base64" not in json.dumps(event_data)


def test_submit_pdf_attachment_passes_native_file_without_persisting_data_url(tmp_path: Path) -> None:
    workspace = _workspace(tmp_path)
    app = create_app(project_space_root=str(tmp_path), no_login=True)
    client = TestClient(app)
    thread_id = client.post("/api/workspaces/default/threads", json={}).json()["thread"]["thread_id"]
    from pypdf import PdfWriter

    pdf_buffer = BytesIO()
    pdf_writer = PdfWriter()
    pdf_writer.add_blank_page(width=72, height=72)
    pdf_writer.write(pdf_buffer)
    pdf_data = base64.b64encode(pdf_buffer.getvalue()).decode("ascii")
    submitted = client.post(
        f"/api/threads/{thread_id}/submit",
        json={
            "text": "inspect this PDF",
            "attachments": [
                {
                    "type": "file",
                    "filename": "paper.pdf",
                    "mime_type": "application/pdf",
                    "data": f"data:application/pdf;base64,{pdf_data}",
                }
            ],
        },
    )

    assert submitted.status_code == 200
    user_message = submitted.json()["message"]
    assert "data:application/pdf" not in str(user_message)
    content = app.state.test_execution.calls[-1]["input"]["messages"][0]["content"]
    assert isinstance(content, list)
    assert content[0]["type"] == "text"
    assert content[1]["type"] == "file"
    assert content[1]["mime_type"] == "application/pdf"
    assert content[1]["base64"] == pdf_data
    multimodal_events = [
        event
        for event in ThreadEventBroker(workspace=workspace).replay(thread_id)
        if event.event == "multimodal.prepared"
    ]
    assert multimodal_events
    event_data = multimodal_events[-1].data
    assert event_data["attachments"][0]["sent_to_model"] is True
    assert event_data["attachments"][0]["sent_as"] == "file"
    assert "base64" not in json.dumps(event_data)


def test_submit_docx_attachments_use_native_only_when_bounded_preflight_fits(
    tmp_path: Path,
) -> None:
    workspace = _workspace(tmp_path)
    app = create_app(project_space_root=str(tmp_path), no_login=True)
    client = TestClient(app)
    thread_id = client.post("/api/workspaces/default/threads", json={}).json()["thread"]["thread_id"]
    from docx import Document

    office_buffer = BytesIO()
    office_document = Document()
    office_document.add_paragraph("compact report")
    office_document.save(office_buffer)
    office_bytes = base64.b64encode(office_buffer.getvalue()).decode("ascii")
    large_office_buffer = BytesIO()
    large_office_document = Document()
    large_office_document.add_paragraph("large document text " * 8_000)
    large_office_document.save(large_office_buffer)
    large_office_bytes = base64.b64encode(large_office_buffer.getvalue()).decode("ascii")
    mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    submitted = client.post(
        f"/api/threads/{thread_id}/submit",
        json={
            "text": "inspect this report",
            "attachments": [
                {
                    "type": "file",
                    "filename": "report.docx",
                    "mime_type": mime,
                    "data": f"data:{mime};base64,{office_bytes}",
                },
                {
                    "type": "file",
                    "filename": "large-report.docx",
                    "mime_type": mime,
                    "data": f"data:{mime};base64,{large_office_bytes}",
                },
            ],
        },
    )

    assert submitted.status_code == 200
    content = app.state.test_execution.calls[-1]["input"]["messages"][0]["content"]
    assert isinstance(content, list)
    assert content[1]["type"] == "file"
    assert content[1]["mime_type"] == mime
    assert content[1]["base64"] == office_bytes
    assert large_office_bytes not in json.dumps(content)
    events = [
        event
        for event in ThreadEventBroker(workspace=workspace).replay(thread_id)
        if event.event == "multimodal.prepared"
    ]
    assert events
    attachments = events[-1].data["attachments"]
    assert attachments[0]["sent_to_model"] is True
    assert attachments[1]["sent_to_model"] is False
    assert "bounded read_file pagination" in attachments[1]["warnings"][0]


def awaitable_result(awaitable):
    import asyncio

    async def _run():
        result = await awaitable
        await asyncio.sleep(0)
        return result

    return asyncio.run(_run())


def test_message_part_pages_keep_ref_and_todos_use_full_current_turn(
    tmp_path: Path,
) -> None:
    workspace = _workspace(tmp_path)
    app = create_app(project_space_root=str(tmp_path), no_login=True)
    client = TestClient(app)
    thread_id = client.post(
        "/api/workspaces/default/threads",
        json={"title": "Long turn"},
    ).json()["thread"]["thread_id"]
    store = ThreadStore(workspace=workspace, workspace_id="default")
    store.append_message(
        ThreadMessage(
            id="msg_user_long_parts",
            thread_id=thread_id,
            role="user",
            status="completed",
            parts=[
                MessagePart(
                    id="part_user_long_parts",
                    type="text",
                    text="Run the bounded review.",
                    status="completed",
                )
            ],
        )
    )
    parts = [
        MessagePart(
            id=f"part_long_{index:03d}",
            type="reasoning",
            text=f"step {index}",
            status="completed",
        )
        for index in range(85)
    ]
    parts[10] = MessagePart(
        id="part_todo_initial",
        type="tool-call",
        status="completed",
        meta={
            "tool": "write_todos",
            "agent_name": "litreview_agent",
            "input": {
                "todos": [
                    {"content": "Search representative work", "status": "completed"},
                    {"content": "Write synthesis", "status": "in_progress"},
                ]
            },
        },
    )
    parts[84] = MessagePart(
        id="part_todo_final",
        type="tool-call",
        status="completed",
        meta={
            "tool": "write_todos",
            "agent_name": "litreview_agent",
            "input": {
                "todos": [
                    {"content": "Search representative work", "status": "completed"},
                    {"content": "Write synthesis", "status": "completed"},
                ]
            },
        },
    )
    store.append_message(
        ThreadMessage(
            id="msg_assistant_long_parts",
            thread_id=thread_id,
            role="assistant",
            status="completed",
            parts=parts,
        )
    )

    message_page = client.get(f"/api/threads/{thread_id}/messages").json()
    assistant = message_page["messages"][-1]
    assert assistant["parts_page"]["shown_count"] == 20
    assert assistant["parts_page"]["total_count"] == 85
    assert message_page["todo_parts"][0]["summary"] == "2 of 2 items complete."
    assert message_page["todo_parts"][0]["id"] == "part_todo_final"

    ref = assistant["parts_page"]["full_content_ref"]
    cursor = assistant["parts_page"]["next_cursor"]
    shown_counts = [20]
    while cursor:
        response = client.get(ref, params={"cursor": cursor, "limit": 20})
        assert response.status_code == 200
        page = response.json()["page"]
        shown_counts.append(page["shown_count"])
        cursor = page["next_cursor"]
        if cursor:
            assert page["full_content_ref"] == ref
    assert shown_counts == [20, 40, 60, 80, 85]


def test_completed_message_preserves_actual_canonical_todo_states(
    tmp_path: Path,
) -> None:
    workspace = _workspace(tmp_path)
    app = create_app(project_space_root=str(tmp_path), no_login=True)
    client = TestClient(app)
    thread_id = client.post(
        "/api/workspaces/default/threads",
        json={"title": "Terminal todo"},
    ).json()["thread"]["thread_id"]
    store = ThreadStore(workspace=workspace, workspace_id="default")
    store.append_message(
        ThreadMessage(
            id="msg_user_terminal_todo",
            thread_id=thread_id,
            role="user",
            status="completed",
            parts=[MessagePart(id="part_user_terminal_todo", type="text", text="Run it.", status="completed")],
        )
    )
    assistant = ThreadMessage(
        id="msg_assistant_terminal_todo",
        thread_id=thread_id,
        role="assistant",
        status="completed",
        parts=[
            MessagePart(
                id="part_root_todo",
                type="tool-call",
                status="completed",
                meta={
                    "tool": "write_todos",
                    "agent_name": "litreview_agent",
                    "input": {"todos": [{"content": "Write synthesis", "status": "completed"}]},
                },
            ),
            MessagePart(
                id="part_child_todo",
                type="tool-call",
                status="completed",
                meta={
                    "tool": "write_todos",
                    "agent_name": "general-purpose",
                    "input": {
                        "todos": [
                            {"content": "Collect papers", "status": "completed"},
                            {"content": "Return handoff", "status": "in_progress"},
                        ]
                    },
                },
            ),
        ],
    )
    store.append_message(assistant)

    message_page = client.get(f"/api/threads/{thread_id}/messages").json()
    assert [part["id"] for part in message_page["todo_parts"]] == ["part_child_todo", "part_root_todo"]

    event = ThreadEventBroker(workspace=workspace).emit(
        thread_id,
        "message.completed",
        message_id=assistant.id,
        status="completed",
        data={"message": assistant.model_dump(mode="json")},
    )
    projected = project_event(event, workspace=workspace)
    assert [part.id for part in projected.data.todo_parts] == ["part_child_todo", "part_root_todo"]


def test_message_page_projects_active_tool_outside_inline_part_page(
    tmp_path: Path,
) -> None:
    workspace = _workspace(tmp_path)
    client = TestClient(create_app(project_space_root=str(tmp_path), no_login=True))
    thread_id = client.post(
        "/api/workspaces/default/threads",
        json={"title": "Long active tool"},
    ).json()["thread"]["thread_id"]
    store = ThreadStore(workspace=workspace, workspace_id="default")
    store.append_message(
        ThreadMessage(
            id="msg_user_active_tool",
            thread_id=thread_id,
            role="user",
            status="completed",
            parts=[MessagePart(id="part_user_active_tool", type="text", text="Run it.", status="completed")],
        )
    )
    parts = [
        MessagePart(
            id=f"part_done_{index}",
            type="tool-call",
            status="completed",
            meta={"tool": "read_file", "input": {"file_path": f"notes/{index}.md"}},
        )
        for index in range(30)
    ]
    parts.append(
        MessagePart(
            id="part_remote_active",
            type="tool-call",
            status="running",
            meta={
                "tool": "remote_submission",
                "input": {"task": "UMA-S stability screening"},
                "agent_name": "materials_worker",
                "started_at": 1_700_000_000.0,
            },
        )
    )
    store.append_message(
        ThreadMessage(
            id="msg_assistant_active_tool",
            thread_id=thread_id,
            role="assistant",
            status="streaming",
            parts=parts,
        )
    )

    payload = client.get(f"/api/threads/{thread_id}/messages").json()

    assert payload["messages"][-1]["parts_page"]["shown_count"] == 20
    assert [part["id"] for part in payload["active_parts"]] == ["part_remote_active"]
    active = payload["active_parts"][0]
    assert active["title"] == "Materials · Remote calculation"
    assert active["started_at"] == 1_700_000_000.0
    assert active["summary"] == "Remote calculation is executing; waiting for its terminal result."


def test_notify_progress_projects_as_semantic_progress_not_active_tool(
    tmp_path: Path,
) -> None:
    projected = project_tool_part(
        {
            "id": "part_notify",
            "type": "tool-call",
            "status": "running",
            "meta": {
                "tool": "notify_progress",
                "input": {
                    "summary": "Literature screening is complete; candidate design is starting.",
                    "next_step": "Delegate three bounded oxide-design branches.",
                },
            },
        },
        workspace=tmp_path,
        thread_id="thread_notify",
        message_id="msg_notify",
    )

    assert projected.type == "progress"
    assert projected.status == "completed"
    assert projected.summary.startswith("Literature screening is complete")
    assert "**Next:** Delegate three bounded oxide-design branches." in projected.text


def test_thread_source_enrichment_replaces_opaque_namespace(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        server,
        "_tool_source_index",
        lambda _run_dir: {
            "write_todos": [
                {
                    "agent_name": "materials_worker",
                    "subagent_source": "materials_worker",
                    "input_key": server._canonical_json(
                        {"todos": [{"content": "Prepare input", "status": "in_progress"}]}
                    ),
                }
            ]
        },
    )
    messages = [
        {
            "meta": {"run_id": "run_opaque_source"},
            "parts": [
                {
                    "type": "tool-call",
                    "meta": {
                        "tool": "write_todos",
                        "agent_name": "tools:da002fa7-6e8f-a97a-2045-fd9eb51d8b06",
                        "input": {"todos": [{"content": "Prepare input", "status": "in_progress"}]},
                    },
                }
            ],
        }
    ]

    enriched = server._enrich_thread_message_tool_sources(messages, workspace=tmp_path)

    part_meta = enriched[0]["parts"][0]["meta"]
    assert part_meta["agent_name"] == "materials_worker"
    assert part_meta["subagent_source"] == "materials_worker"
