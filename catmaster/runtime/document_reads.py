from __future__ import annotations

import asyncio
import tempfile
import threading
from collections import OrderedDict
from contextlib import contextmanager
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from typing import Any, BinaryIO, Iterator
from zipfile import BadZipFile, ZipFile

from langchain.agents.middleware import AgentMiddleware
from langchain_core.messages import ToolMessage

from catmaster.tools.base import normalize_workspace_virtual_path


SUPPORTED_DOCUMENT_SUFFIXES = frozenset(
    {".pdf", ".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx"}
)
PAGINATED_DOCUMENT_SUFFIXES = frozenset({".pdf", ".docx", ".xlsx", ".pptx"})

# Native file blocks are useful for compact documents, but DeepAgents 0.7.4
# reads every binary file in full and ignores read_file offset/limit. Keep the
# native lane well below OpenAI's 50 MB per-request file ceiling because base64
# expands the HTTP payload and the provider's extracted document text still
# consumes model context.
MAX_NATIVE_DOCUMENT_FILE_BYTES = 8 * 1024 * 1024
MAX_NATIVE_DOCUMENT_TEXT_CHARS = 60_000
MAX_NATIVE_DOCUMENT_LINES = 100
MAX_NATIVE_PDF_PAGES = 20

# These are parser safety bounds, not agent-visible validation metadata.
MAX_DOCUMENT_FILE_BYTES = 50 * 1024 * 1024
MAX_OFFICE_ARCHIVE_MEMBERS = 10_000
MAX_OFFICE_UNCOMPRESSED_BYTES = 256 * 1024 * 1024
MAX_OFFICE_MEMBER_BYTES = 128 * 1024 * 1024
MAX_DOCUMENT_LINES_PER_READ = 200
MAX_DOCUMENT_TEXT_CHARS_PER_READ = 60_000
MAX_CANONICAL_LINE_CHARS = 4_000
MAX_PARSED_DOCUMENT_CACHE_ENTRIES = 4


class DocumentReadError(ValueError):
    """A document cannot be read through the bounded local text view."""


@dataclass(frozen=True, slots=True)
class NativeDocumentDecision:
    allowed: bool
    reason: str = ""


@dataclass(frozen=True, slots=True)
class BoundedDocumentPage:
    content: str
    next_offset: int
    returned_lines: int


class _ParsedDocumentLines:
    """Incrementally parse one document into a bounded, process-local spool."""

    def __init__(self, path: Path, *, mtime_ns: int, size: int) -> None:
        self.path = path
        self.mtime_ns = mtime_ns
        self.size = size
        self.in_use = 0
        self.retired = False
        self._iterator = _iter_document_lines(path)
        self._store: BinaryIO = tempfile.TemporaryFile(mode="w+b")
        self._positions = [0]
        self._complete = False
        self._failure: DocumentReadError | None = None
        self._closed = False
        self._lock = threading.RLock()

    def matches(self, *, mtime_ns: int, size: int) -> bool:
        return self.mtime_ns == mtime_ns and self.size == size

    def ensure_line(self, index: int) -> bool:
        """Parse through *index* and report whether that line exists."""

        with self._lock:
            if self._failure is not None:
                raise self._failure
            while len(self._positions) - 1 <= index and not self._complete:
                try:
                    line = next(self._iterator)
                except StopIteration:
                    self._complete = True
                    break
                except DocumentReadError as exc:
                    self._failure = exc
                    self._complete = True
                    raise
                payload = line.encode("utf-8")
                self._store.seek(0, 2)
                self._store.write(payload)
                self._positions.append(self._positions[-1] + len(payload))
            return index < len(self._positions) - 1

    def read_line(self, index: int) -> str:
        with self._lock:
            if index < 0 or index >= len(self._positions) - 1:
                raise IndexError(index)
            start = self._positions[index]
            length = self._positions[index + 1] - start
            self._store.seek(start)
            return self._store.read(length).decode("utf-8")

    def close(self) -> None:
        with self._lock:
            if self._closed:
                return
            self._closed = True
            close_iterator = getattr(self._iterator, "close", None)
            if callable(close_iterator):
                close_iterator()
            self._store.close()


_PARSED_DOCUMENT_CACHE_LOCK = threading.RLock()
_PARSED_DOCUMENT_CACHE: OrderedDict[Path, _ParsedDocumentLines] = OrderedDict()


def _retire_parsed_document_entry(entry: _ParsedDocumentLines) -> None:
    entry.retired = True
    if entry.in_use == 0:
        entry.close()


def _evict_parsed_document_entries() -> None:
    while len(_PARSED_DOCUMENT_CACHE) > MAX_PARSED_DOCUMENT_CACHE_ENTRIES:
        victim_path = next(
            (
                path
                for path, entry in _PARSED_DOCUMENT_CACHE.items()
                if entry.in_use == 0
            ),
            None,
        )
        if victim_path is None:
            return
        victim = _PARSED_DOCUMENT_CACHE.pop(victim_path)
        _retire_parsed_document_entry(victim)


@contextmanager
def _parsed_document_lines(path: Path) -> Iterator[_ParsedDocumentLines]:
    try:
        stat = path.stat()
    except OSError as exc:
        raise DocumentReadError(f"could not inspect document: {exc}") from exc

    with _PARSED_DOCUMENT_CACHE_LOCK:
        entry = _PARSED_DOCUMENT_CACHE.get(path)
        if entry is not None and not entry.matches(
            mtime_ns=stat.st_mtime_ns,
            size=stat.st_size,
        ):
            _PARSED_DOCUMENT_CACHE.pop(path)
            _retire_parsed_document_entry(entry)
            entry = None
        if entry is None:
            entry = _ParsedDocumentLines(
                path,
                mtime_ns=stat.st_mtime_ns,
                size=stat.st_size,
            )
            _PARSED_DOCUMENT_CACHE[path] = entry
        else:
            _PARSED_DOCUMENT_CACHE.move_to_end(path)
        entry.in_use += 1
        _evict_parsed_document_entries()

    try:
        yield entry
    finally:
        with _PARSED_DOCUMENT_CACHE_LOCK:
            entry.in_use -= 1
            if entry.retired and entry.in_use == 0:
                entry.close()
            _evict_parsed_document_entries()


def _resolve_virtual_document(files_root: Path, file_path: str) -> Path:
    root = Path(files_root).expanduser().resolve()
    raw = normalize_workspace_virtual_path(
        str(file_path or "").strip().replace("\\", "/")
    )
    virtual = PurePosixPath(raw)
    if not raw.startswith("/") or ".." in virtual.parts or raw.startswith("~"):
        raise DocumentReadError(
            "file_path must be an absolute virtual workspace path without traversal"
        )
    source = (root / raw.lstrip("/")).resolve()
    try:
        source.relative_to(root)
    except ValueError as exc:
        raise DocumentReadError("document path is outside the workspace") from exc
    if not source.is_file():
        raise DocumentReadError(f"document file not found: {raw}")
    if source.suffix.lower() not in SUPPORTED_DOCUMENT_SUFFIXES:
        raise DocumentReadError(f"unsupported document type: {source.suffix or 'none'}")
    return source


def _inspect_document(path: Path) -> int:
    try:
        size = path.stat().st_size
    except OSError as exc:
        raise DocumentReadError(f"could not inspect document: {exc}") from exc
    if size > MAX_DOCUMENT_FILE_BYTES:
        raise DocumentReadError(
            "document exceeds the bounded local reading limit; split it into smaller source files"
        )
    return size


def _validate_office_archive(path: Path) -> None:
    try:
        with ZipFile(path) as archive:
            members = archive.infolist()
    except (BadZipFile, OSError) as exc:
        raise DocumentReadError(f"invalid Office document: {exc}") from exc
    if len(members) > MAX_OFFICE_ARCHIVE_MEMBERS:
        raise DocumentReadError("Office document contains too many internal files")
    expanded = 0
    for member in members:
        if member.flag_bits & 0x1:
            raise DocumentReadError("encrypted Office documents are not supported")
        if member.file_size > MAX_OFFICE_MEMBER_BYTES:
            raise DocumentReadError("Office document contains an oversized internal file")
        expanded += member.file_size
        if expanded > MAX_OFFICE_UNCOMPRESSED_BYTES:
            raise DocumentReadError("Office document expands beyond the local parser limit")


def _split_canonical_line(value: Any) -> Iterator[str]:
    text = str(value or "").replace("\r\n", "\n").replace("\r", "\n")
    for raw_line in text.split("\n"):
        line = raw_line.strip()
        if not line:
            continue
        for start in range(0, len(line), MAX_CANONICAL_LINE_CHARS):
            yield line[start : start + MAX_CANONICAL_LINE_CHARS]


def _iter_pdf_lines(path: Path) -> Iterator[str]:
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path), strict=False)
        for page_index, page in enumerate(reader.pages, start=1):
            yield f"[PDF page {page_index}]"
            try:
                text = page.extract_text() or ""
            except Exception as exc:
                yield f"[Text extraction unavailable on page {page_index}: {exc}]"
                continue
            yield from _split_canonical_line(text)
    except DocumentReadError:
        raise
    except Exception as exc:
        raise DocumentReadError(f"could not parse PDF: {exc}") from exc


def _iter_docx_lines(path: Path) -> Iterator[str]:
    _validate_office_archive(path)
    try:
        from docx import Document
        from docx.table import Table

        document = Document(str(path))
        for block in document.iter_inner_content():
            if isinstance(block, Table):
                for row in block.rows:
                    values = [" ".join(cell.text.split()) for cell in row.cells]
                    yield from _split_canonical_line("\t".join(values))
            else:
                yield from _split_canonical_line(getattr(block, "text", ""))
    except DocumentReadError:
        raise
    except Exception as exc:
        raise DocumentReadError(f"could not parse DOCX: {exc}") from exc


def _iter_xlsx_lines(path: Path) -> Iterator[str]:
    _validate_office_archive(path)
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(
            filename=path,
            read_only=True,
            data_only=True,
        )
        try:
            for sheet in workbook.worksheets:
                yield f"[XLSX sheet: {sheet.title}]"
                for row in sheet.iter_rows(values_only=True):
                    values = ["" if value is None else str(value) for value in row]
                    while values and not values[-1]:
                        values.pop()
                    if values:
                        yield from _split_canonical_line("\t".join(values))
        finally:
            workbook.close()
    except DocumentReadError:
        raise
    except Exception as exc:
        raise DocumentReadError(f"could not parse XLSX: {exc}") from exc


def _iter_pptx_shape_lines(shape: Any) -> Iterator[str]:
    child_shapes = getattr(shape, "shapes", None)
    if child_shapes is not None:
        for child in child_shapes:
            yield from _iter_pptx_shape_lines(child)
    table = getattr(shape, "table", None) if getattr(shape, "has_table", False) else None
    if table is not None:
        for row in table.rows:
            yield from _split_canonical_line(
                "\t".join(" ".join(cell.text.split()) for cell in row.cells)
            )
        return
    if getattr(shape, "has_text_frame", False):
        yield from _split_canonical_line(getattr(shape, "text", ""))


def _iter_pptx_lines(path: Path) -> Iterator[str]:
    _validate_office_archive(path)
    try:
        from pptx import Presentation

        presentation = Presentation(str(path))
        for slide_index, slide in enumerate(presentation.slides, start=1):
            yield f"[PPTX slide {slide_index}]"
            for shape in slide.shapes:
                yield from _iter_pptx_shape_lines(shape)
    except DocumentReadError:
        raise
    except Exception as exc:
        raise DocumentReadError(f"could not parse PPTX: {exc}") from exc


def _iter_document_lines(path: Path) -> Iterator[str]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        yield from _iter_pdf_lines(path)
        return
    if suffix == ".docx":
        yield from _iter_docx_lines(path)
        return
    if suffix == ".xlsx":
        yield from _iter_xlsx_lines(path)
        return
    if suffix == ".pptx":
        yield from _iter_pptx_lines(path)
        return
    raise DocumentReadError(
        f"{suffix[1:].upper()} cannot be paginated locally; convert it to the corresponding OOXML format"
    )


def _pdf_page_count(path: Path) -> int:
    try:
        from pypdf import PdfReader

        return len(PdfReader(str(path), strict=False).pages)
    except Exception as exc:
        raise DocumentReadError(f"could not inspect PDF pages: {exc}") from exc


def native_document_input_decision(path: Path) -> NativeDocumentDecision:
    source = Path(path).expanduser().resolve()
    suffix = source.suffix.lower()
    if suffix not in SUPPORTED_DOCUMENT_SUFFIXES or not source.is_file():
        return NativeDocumentDecision(False, "not a readable supported document")
    try:
        size = _inspect_document(source)
        if size > MAX_NATIVE_DOCUMENT_FILE_BYTES:
            return NativeDocumentDecision(False, "raw file is too large for one native request")
        if suffix not in PAGINATED_DOCUMENT_SUFFIXES:
            return NativeDocumentDecision(True)
        if suffix == ".pdf" and _pdf_page_count(source) > MAX_NATIVE_PDF_PAGES:
            return NativeDocumentDecision(False, "PDF has too many pages for one native request")

        chars = 0
        line_index = 0
        with _parsed_document_lines(source) as parsed:
            while parsed.ensure_line(line_index):
                line = parsed.read_line(line_index)
                line_index += 1
                chars += len(line)
                if (
                    line_index > MAX_NATIVE_DOCUMENT_LINES
                    or chars > MAX_NATIVE_DOCUMENT_TEXT_CHARS
                ):
                    return NativeDocumentDecision(
                        False,
                        "document text is too large for one native request",
                    )
        return NativeDocumentDecision(True)
    except DocumentReadError as exc:
        return NativeDocumentDecision(False, str(exc))


def read_bounded_document(
    files_root: Path,
    *,
    file_path: str,
    offset: int = 0,
    limit: int = 100,
) -> BoundedDocumentPage:
    source = _resolve_virtual_document(files_root, file_path)
    _inspect_document(source)
    if source.suffix.lower() not in PAGINATED_DOCUMENT_SUFFIXES:
        raise DocumentReadError(
            f"{source.suffix[1:].upper()} has no bounded local reader; convert it to DOCX, XLSX, or PPTX first"
        )
    start = max(0, int(offset))
    requested_limit = int(limit)
    if requested_limit <= 0:
        raise DocumentReadError("limit must be greater than zero for document reads")
    effective_limit = min(requested_limit, MAX_DOCUMENT_LINES_PER_READ)

    selected: list[str] = []
    selected_chars = 0
    has_more = False
    with _parsed_document_lines(source) as parsed:
        index = start
        while parsed.ensure_line(index):
            line = parsed.read_line(index)
            if len(selected) >= effective_limit:
                has_more = True
                break
            addition = len(line) + (1 if selected else 0)
            if selected and selected_chars + addition > MAX_DOCUMENT_TEXT_CHARS_PER_READ:
                has_more = True
                break
            selected.append(line)
            selected_chars += addition
            index += 1

    if not selected:
        return BoundedDocumentPage(
            content=f"Document text view `{file_path}` has no lines at offset {start}.",
            next_offset=0,
            returned_lines=0,
        )

    numbered = "\n".join(
        f"{start + index + 1:>6}  {line}"
        for index, line in enumerate(selected)
    )
    next_offset = start + len(selected) if has_more else 0
    content = (
        f"Bounded {source.suffix[1:].upper()} text view: `{file_path}`\n"
        f"Lines {start + 1}-{start + len(selected)}\n\n{numbered}"
    )
    if next_offset:
        content += (
            "\n\nMore document text remains. Continue with "
            f"`read_file(file_path={file_path!r}, offset={next_offset}, limit={requested_limit})`."
        )
    return BoundedDocumentPage(
        content=content,
        next_offset=next_offset,
        returned_lines=len(selected),
    )


class BoundedDocumentReadMiddleware(AgentMiddleware):
    """Keep DeepAgents `read_file` while bounding documents that do not fit one request."""

    def __init__(self, *, files_root: Path) -> None:
        self.files_root = Path(files_root).expanduser().resolve()

    @property
    def name(self) -> str:
        return "catmaster_bounded_document_reads"

    @staticmethod
    def _tool_call(request: Any) -> tuple[str, str, dict[str, Any]]:
        tool_call = getattr(request, "tool_call", None)
        if not isinstance(tool_call, dict):
            return "", "", {}
        args = tool_call.get("args")
        return (
            str(tool_call.get("name") or "").strip(),
            str(tool_call.get("id") or "").strip(),
            dict(args) if isinstance(args, dict) else {},
        )

    def _bounded_result(self, request: Any) -> ToolMessage | None:
        tool_name, tool_call_id, args = self._tool_call(request)
        file_path = str(args.get("file_path") or "").strip()
        if tool_name != "read_file" or Path(file_path).suffix.lower() not in SUPPORTED_DOCUMENT_SUFFIXES:
            return None
        try:
            offset = int(args.get("offset", 0))
            limit = int(args.get("limit", 100))
            source = _resolve_virtual_document(self.files_root, file_path)
            decision = native_document_input_decision(source)
            # A default read of a compact document keeps DeepAgents' provider-native
            # file block. Any explicit pagination request, or a document that does
            # not fit the bounded preflight, uses the local text view instead.
            if decision.allowed and offset == 0 and limit == 100:
                return None
            page = read_bounded_document(
                self.files_root,
                file_path=file_path,
                offset=offset,
                limit=limit,
            )
            return ToolMessage(
                content=page.content,
                name="read_file",
                tool_call_id=tool_call_id or "read_file_document",
                status="success",
                additional_kwargs={
                    "read_file_path": file_path,
                    "catmaster_bounded_document_read": True,
                },
            )
        except (DocumentReadError, TypeError, ValueError) as exc:
            return ToolMessage(
                content=f"Error reading document `{file_path}`: {exc}",
                name="read_file",
                tool_call_id=tool_call_id or "read_file_document",
                status="error",
            )

    def wrap_tool_call(self, request: Any, handler: Any) -> Any:
        bounded = self._bounded_result(request)
        return bounded if bounded is not None else handler(request)

    async def awrap_tool_call(self, request: Any, handler: Any) -> Any:
        bounded = await asyncio.to_thread(self._bounded_result, request)
        if bounded is not None:
            return bounded
        return await handler(request)


__all__ = [
    "BoundedDocumentPage",
    "BoundedDocumentReadMiddleware",
    "DocumentReadError",
    "MAX_NATIVE_DOCUMENT_FILE_BYTES",
    "NativeDocumentDecision",
    "PAGINATED_DOCUMENT_SUFFIXES",
    "SUPPORTED_DOCUMENT_SUFFIXES",
    "native_document_input_decision",
    "read_bounded_document",
]
